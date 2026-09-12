"""Excel/CSV → Markdown 轉換模組：把 .xlsx/.xlsm/.csv 轉成「每列一個 ### 區塊」的 Markdown，
輸出慣例與 flatten_md_tables.py 一致（每列自成一個檢索單元，切塊行為與扁平化後的 PDF 表格相同）。

用法：
    python sheet_to_md.py <file...>    # 把轉出的 Markdown 印到 stdout（方便測試單檔）

轉換規則：
- xlsx/xlsm：openpyxl load_workbook(path, data_only=True)（非 read_only 模式，否則
  merged_cells.ranges 不可用）；先掃 merged ranges 建展開映射，再逐列取值展開填值
- csv：編碼鏈 utf-8-sig → utf-16 → big5；分隔符號 csv.Sniffer（限 ,、\\t、;），失敗 fallback ,
- 每個 sheet 輸出 `# {filename} — Sheet: {name}`（CSV 為 `# {filename}`；隱藏 sheet 加 ` (hidden)`）
- 表頭判定：前 10 個非空列中，第一列「≥2 個非空 cell 且全部 ≤40 字元」者為表頭（門檻同
  flatten 的 LABEL_MAX）；其上的非空列保留為純文字 preamble；找不到 → headerless
- 每資料列一個 ### 區塊（鏡射 flatten_md_tables.render_row）：標題 = 首欄值（空則取首個非空
  cell；>60 字元截斷加 "…"，全文以內文行保留）；其餘非空 cell 輸出 `[欄位標籤]: 值`（headerless
  無前綴）；超出表頭欄數的 cell 以無前綴純文字行輸出；無表頭表格不加任何自訂標籤
- 值 → 文字：None 跳過、datetime/date/time → ISO 格式、bool → TRUE/FALSE、整數值的 float
  去小數（2.0 → 2）、cell 內換行換成空格；全空白行跳過、列尾空欄 trim

已知限制：.xls/.ods 不支援（請先另存 .xlsx）；公式取 Excel 存檔的快取值（从未由 Excel 計算
存檔者為空，review 時補）；同 sheet 多表格只用第一個偵測到的表頭；數字顯示原始數值（不含
貨幣/百分比顯示格式）。
"""

import argparse
import csv
import datetime
import io
import os
import sys

TITLE_MAX = 60  # 區塊標題最長字元數，超過則截斷加 "…"（同 flatten_md_tables）
LABEL_MAX = 40  # 表頭 cell「短標籤」判定上限（同 flatten_md_tables）
HEADER_SCAN_ROWS = 10  # 表頭判定只看前 N 個非空列
CSV_ENCODINGS = ("utf-8-sig", "utf-16", "big5")
CSV_SNIFF_DELIMITERS = ",\t;"


def render_value(value):
    """cell 值 → 輸出文字：datetime/date/time → ISO、bool → TRUE/FALSE、整數值的 float 去小數、換行換空格"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def build_merged_map(ws):
    """{(row, col): anchor 值} — merged range 涵蓋的每格填入左上角 anchor 的值（每列帶完整 context）"""
    merged = {}
    for rng in ws.merged_cells.ranges:
        anchor = ws.cell(rng.min_row, rng.min_col).value
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                if (row, col) != (rng.min_row, rng.min_col):
                    merged[(row, col)] = anchor
    return merged


def sheet_rows(ws):
    """worksheet → 字串列（merged cell 已展開、值已轉文字、列尾空欄 trim）"""
    merged = build_merged_map(ws)
    rows = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        cells = [render_value(merged.get((row_idx, col_idx), value)) for col_idx, value in enumerate(row, 1)]
        while cells and not cells[-1]:
            cells.pop()
        rows.append(cells)
    return rows


def read_csv_rows(path):
    """CSV → 字串列（編碼鏈 utf-8-sig → utf-16 → big5；分隔符號 Sniffer；視為單一 sheet）"""
    text = None
    for encoding in CSV_ENCODINGS:
        try:
            with open(path, "r", encoding=encoding, newline="") as f:
                text = f.read()
            break
        except UnicodeError:
            continue
    if text is None:
        raise ValueError(f"could not decode CSV with any of {', '.join(CSV_ENCODINGS)}")
    delimiter = ","
    try:
        delimiter = csv.Sniffer().sniff(text[:65536], delimiters=CSV_SNIFF_DELIMITERS).delimiter
    except csv.Error:
        pass
    rows = []
    for raw in csv.reader(io.StringIO(text), delimiter=delimiter):
        cells = [render_value(cell) for cell in raw]
        while cells and not cells[-1]:
            cells.pop()
        rows.append(cells)
    return rows


def is_header_row(cells):
    """≥2 個非空 cell 且全部 ≤LABEL_MAX 短標籤 → 表頭候選（同 flatten 的 is_short_label_row）"""
    non_empty = [c for c in cells if c]
    return len(non_empty) >= 2 and all(len(c) <= LABEL_MAX for c in non_empty)


def make_title(text):
    return text if len(text) <= TITLE_MAX else text[:TITLE_MAX].rstrip() + "…"


def render_data_row(cells, labels):
    """資料列 → '### 標題' 區塊（有表頭 → [標籤]: 行；無表頭 → 無前綴純文字行；鏡射 flatten_md_tables.render_row）"""
    first = cells[0] if cells and cells[0] else next((c for c in cells[1:] if c), "")
    body = []
    if len(first) > TITLE_MAX:
        # 標題有截斷 → 首 cell 全文以內文行保留，避免內容遺失
        body.append(f"[{labels[0]}]: {first}" if labels and labels[0] else first)
    for i in range(1, len(cells)):
        if not cells[i]:
            continue
        body.append(f"[{labels[i]}]: {cells[i]}" if labels and i < len(labels) and labels[i] else cells[i])
    return "\n\n".join([f"### {make_title(first)}"] + body)


def render_sheet(filename, sheet_name, rows, hidden=False):
    """單一 sheet（或 CSV 整檔）→ Markdown 段落；全空回 None（呼叫端跳過）。sheet_name=None 表示 CSV。"""
    non_empty = [r for r in rows if any(r)]
    if not non_empty:
        return None

    heading = f"# {filename}" if sheet_name is None else f"# {filename} — Sheet: {sheet_name}"
    if hidden:
        heading += " (hidden)"

    # 表頭判定：前 HEADER_SCAN_ROWS 個非空列中第一個短標籤列；其上非空列 = preamble
    labels, preamble, data = None, [], non_empty
    for i, row in enumerate(non_empty[:HEADER_SCAN_ROWS]):
        if is_header_row(row):
            labels = row
            preamble, data = non_empty[:i], non_empty[i + 1:]
            break

    parts = [heading]
    if preamble:
        parts.append("\n".join(" ".join(c for c in row if c) for row in preamble))
    parts.extend(render_data_row(row, labels) for row in data)
    return "\n\n".join(parts)


def convert_file(path):
    """讀取 .xlsx/.xlsm/.csv → 完整 Markdown 字串；全部 sheet 皆空時 raise ValueError（呼叫端轉 per-file error）"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        md = render_sheet(os.path.basename(path), None, read_csv_rows(path))
        if md is None:
            raise ValueError("no data found")
        return md + "\n"
    if ext in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True)
        parts = []
        for ws in wb.worksheets:
            md = render_sheet(os.path.basename(path), ws.title, sheet_rows(ws),
                              hidden=ws.sheet_state != "visible")
            if md is None:
                print(f"  Sheet '{ws.title}' is empty, skipped")
                continue
            parts.append(md)
        if not parts:
            raise ValueError("no data found")
        return "\n\n".join(parts) + "\n"
    if ext == ".xls":
        raise ValueError(".xls is not supported (old binary format); open it in Excel and save as .xlsx first")
    raise ValueError(f"unsupported file type '{ext}' (supported: .xlsx, .xlsm, .csv)")


def main():
    parser = argparse.ArgumentParser(
        description="把 .xlsx/.xlsm/.csv 轉成每列一個 ### 區塊的 Markdown，印到 stdout"
    )
    parser.add_argument("files", nargs="+", help=".xlsx/.xlsm/.csv 檔案")
    args = parser.parse_args()
    failed = False
    for path in args.files:
        try:
            sys.stdout.write(convert_file(path))
        except Exception as e:
            failed = True
            print(f"Error: {path}: {e}", file=sys.stderr)
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
